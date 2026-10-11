from fastapi import FastAPI, Depends, HTTPException
from pydantic import BaseModel
from sqlalchemy import Column, Integer, String, func, text
from sqlalchemy.orm import Session
from fastapi.middleware.cors import CORSMiddleware
from datetime import datetime
from zoneinfo import ZoneInfo
import random
import os
from io import BytesIO
from fastapi.security import HTTPBasic, HTTPBasicCredentials
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
import secrets

from database import Base, engine, get_db, SessionLocal


app = FastAPI()


app.add_middleware( CORSMiddleware, allow_origins=[ "http://localhost:5500", "http://127.0.0.1:5500" ], allow_credentials=True, allow_methods=["*"], allow_headers=["*"], )

FEEDBACK_TYPES = [
    "negative_destructive",
    "negative_supportive",
    "positive_destructive",
    "positive_supportive"
]

class ExperimentResult(Base):
    __tablename__ = "experiment_results"

    id = Column(Integer, primary_key=True, index=True)
    student_id = Column(String, nullable=False)
    stage1_score = Column(Integer, nullable=False)
    stage2_score = Column(Integer, nullable=False)
    feedback_type = Column(String, nullable=False)
    finished_at = Column(String, nullable=True)

class FeedbackCounter(Base):
    __tablename__ = "feedback_counters"

    id = Column(Integer, primary_key=True, index=True)

    feedback_type = Column(String, unique=True, nullable=False)

    count = Column(Integer, nullable=False, default=0)


Base.metadata.create_all(bind=engine)


def initialize_feedback_counters():
    db = SessionLocal()

    try:
        for feedback_type in FEEDBACK_TYPES:
            existing = (
                db.query(FeedbackCounter)
                .filter(
                    FeedbackCounter.feedback_type == feedback_type
                )
                .first()
            )

            if existing is None:
                db.add(
                    FeedbackCounter(
                        feedback_type=feedback_type,
                        count=0
                    )
                )

        db.commit()

    finally:
        db.close()


initialize_feedback_counters()




class ExperimentResultCreate(BaseModel):
    student_id: str
    stage1_score: int
    stage2_score: int
    feedback_type: str





@app.get("/")
def root():
    return {
        "message": "Backend is running"
    }


@app.get("/feedback/next")
def get_next_feedback(
    db: Session = Depends(get_db)
):
    db.execute(text("BEGIN IMMEDIATE"))

    try:
        feedback_counters = db.query(FeedbackCounter).all()

        min_count = min(
            counter.count
            for counter in feedback_counters
        )

        least_used = [
            counter
            for counter in feedback_counters
            if counter.count == min_count
        ]

        selected_counter = random.choice(least_used)

        selected_counter.count += 1

        db.commit()

        return {
            "feedback_type": selected_counter.feedback_type,
            "count": selected_counter.count
        }

    except Exception:
        db.rollback()
        raise


@app.post("/results")
def create_result(
    result: ExperimentResultCreate,
    db: Session = Depends(get_db)
):
    finished_at = datetime.now(ZoneInfo("Asia/Tehran"))

    db_result = ExperimentResult(
        student_id=result.student_id,
        stage1_score=result.stage1_score,
        stage2_score=result.stage2_score,
        feedback_type=result.feedback_type,
        finished_at=datetime.now(
            ZoneInfo("Asia/Tehran")
        ).isoformat(timespec="seconds")
    )

    db.add(db_result)
    db.commit()
    db.refresh(db_result)

    return {
        "message": "Result saved successfully",
        "id": db_result.id,
        #"finished_at": db_result.finished_at
    }

export_security = HTTPBasic()

def verify_export_access(
    credentials: HTTPBasicCredentials = Depends(export_security),
):
    expected_username = os.getenv("EXPORT_USERNAME")
    expected_password = os.getenv("EXPORT_PASSWORD")

    if not expected_username or not expected_password:
        raise HTTPException(
            status_code=503,
            detail="Export credentials are not configured",
        )

    username_ok = secrets.compare_digest(
        credentials.username, expected_username
    )
    password_ok = secrets.compare_digest(
        credentials.password, expected_password
    )

    if not (username_ok and password_ok):
        raise HTTPException(
            status_code=401,
            detail="Invalid username or password",
            headers={"WWW-Authenticate": "Basic"},
        )

@app.get("/results")
def get_results(
    db: Session = Depends(get_db),
    _authorized: None = Depends(verify_export_access),
):
    results = db.query(ExperimentResult).all()

    return [
        {
            "id": result.id,
            "student_id": result.student_id,
            "stage1_score": result.stage1_score,
            "stage2_score": result.stage2_score,
            "feedback_type": result.feedback_type,
            "finished_at": result.finished_at,
        }
        for result in results
    ]

@app.get("/results/export")
def export_results_to_excel(
    db: Session = Depends(get_db),
    _authorized: None = Depends(verify_export_access),
):
    results = (
        db.query(ExperimentResult)
        .order_by(ExperimentResult.id.asc())
        .all()
    )

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Experiment Results"

    headers = [
        "Student ID",
        "Stage 1 Score",
        "Stage 2 Score",
        "Feedback Type",
        "Finished At",
    ]
    sheet.append(headers)

    for result in results:
        sheet.append([
            result.student_id,
            result.stage1_score,
            result.stage2_score,
            result.feedback_type,
            result.finished_at,
        ])

    # قالب‌بندی عنوان ستون‌ها
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(
            fill_type="solid",
            fgColor="1976D2",
        )
        cell.alignment = Alignment(horizontal="center")

    # تنظیم عرض ستون‌ها
    for column_cells in sheet.columns:
        max_length = max(
            len(str(cell.value)) if cell.value is not None else 0
            for cell in column_cells
        )
        column_letter = get_column_letter(column_cells[0].column)
        sheet.column_dimensions[column_letter].width = min(
            max_length + 3, 35
        )

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions

    output = BytesIO()
    workbook.save(output)
    output.seek(0)

    return StreamingResponse(
        output,
        media_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
        headers={
            "Content-Disposition": (
                'attachment; filename="experiment_results.xlsx"'
            )
        },
    )